import json

from openpyxl import load_workbook

from governance_app.database_runtime import database_for
from governance_app.db import connect, initialize_database
from governance_app.exporter import export_city_issue_packages
from governance_app.migrations import apply_migrations
from governance_app.server import create_app


def _json(app, method, path, payload=None):
    response = app.handle_test_request(
        method, path, json.dumps(payload, ensure_ascii=False) if payload is not None else "")
    return response[0], json.loads(response[2])


def test_local_tower_rent_correction_targets_one_record_and_reaudit_uses_it(
    app_config, sample_workbook, tmp_path
):
    workbook = load_workbook(sample_workbook)
    sheet = workbook["铁塔租费台账"]
    sheet.cell(1, 9, "产品单元数")
    sheet.cell(2, 9, 0)
    sheet.cell(1, 10, "需求单号")
    sheet.cell(2, 10, "ORDER-1")
    sheet.append([2, "HZ001", "西湖一站", "杭州", "西湖", "TT002", "第二租费记录", 10000, 0, "ORDER-2"])
    path = tmp_path / "two-rents.xlsx"
    workbook.save(path)
    initialize_database(app_config)
    app = create_app(app_config)
    status, imported = _json(app, "POST", "/api/import", {"path": str(path)})
    assert status == 200
    batch_id = imported["batch_id"]
    _, listing = _json(app, "GET", f"/api/local/tower-rents?batch_id={batch_id}")
    first, second = [row["row_id"] for row in listing["tower_rents"]]
    assert _json(app, "POST", "/api/audit", {"batch_id": batch_id})[0] == 200
    _, before = _json(app, "GET", f"/api/issues?batch_id={batch_id}")
    assert len([issue for issue in before["issues"] if issue["rule_id"] == "tower_product_units_zero_fee_nonzero"]) == 2
    payload = {"batch_id": batch_id,
               "changes": {"产品服务费合计（元/年）（不含税）": 0},
               "evidence": "合同及核实单", "operator": "省公司", "error_cause": "计费错填",
               "source": "市州反馈", "idempotency_key": "rent-1"}
    assert _json(app, "POST", f"/api/local/tower-rents/{first}/corrections", payload) == (200, {"version": 1})
    assert _json(app, "POST", f"/api/local/tower-rents/{first}/corrections", payload) == (200, {"version": 1})
    assert _json(app, "POST", f"/api/local/tower-rents/{first}/corrections",
                 {**payload, "changes": {"产品服务费合计（元/年）（不含税）": 1}})[0] == 409
    _, first_detail = _json(app, "GET", f"/api/local/tower-rents/{first}?batch_id={batch_id}")
    _, second_detail = _json(app, "GET", f"/api/local/tower-rents/{second}?batch_id={batch_id}")
    field = "产品服务费合计（元/年）（不含税）"
    assert first_detail["source"][field] == 10000
    assert first_detail["current"][field] == 0
    assert first_detail["versions"][0]["evidence"] == "合同及核实单"
    assert second_detail["current"][field] == 10000
    assert second_detail["version"] == 0
    assert _json(app, "POST", "/api/audit", {"batch_id": batch_id})[0] == 200
    _, after = _json(app, "GET", f"/api/issues?batch_id={batch_id}")
    active = [issue for issue in after["issues"] if issue["rule_id"] == "tower_product_units_zero_fee_nonzero"
              and issue["status"] != "resolved_by_reaudit"]
    assert len(active) == 1
    assert active[0]["telecom_site_code"] == "HZ001"
    _, rows = _json(app, "GET", f"/api/ledger-rows?batch_id={batch_id}&ledger_type=tower_rent")
    assert [row["raw"][field] for row in rows["rows"]] == [0, 10000]
    exported = export_city_issue_packages(app_config, batch_id)
    assert exported
    sheet = load_workbook(exported[0])["整改问题清单"]
    headings = {cell.value: cell.column for cell in sheet[1]}
    rent_issues = [row for row in sheet.iter_rows(min_row=2, values_only=True)
                   if row[headings["规则编号"] - 1] == "tower_product_units_zero_fee_nonzero"]
    assert len(rent_issues) == 1
    assert rent_issues[0][headings["原始字段值"] - 1] == "10000"
    _, reimported = _json(app, "POST", "/api/import", {"path": str(path)})
    _, newer = _json(app, "GET", f"/api/local/tower-rents?batch_id={reimported['batch_id']}")
    _, carried = _json(app, "GET", f"/api/local/tower-rents/{newer['tower_rents'][0]['row_id']}?batch_id={reimported['batch_id']}")
    assert carried["source"][field] == 10000
    assert carried["current"][field] == 0
    assert carried["version"] == 1


def test_rent_without_agreement_number_is_not_silently_linked_across_batches(
    app_config, sample_workbook
):
    initialize_database(app_config)
    app = create_app(app_config)
    first_batch = _json(app, "POST", "/api/import", {"path": str(sample_workbook)})[1]["batch_id"]
    first_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={first_batch}")[1]["tower_rents"][0]["row_id"]
    payload = {"batch_id": first_batch,
               "changes": {"产品服务费合计（元/年）（不含税）": 5000},
               "evidence": "核实材料", "operator": "省公司", "error_cause": "计费错误",
               "source": "合同", "idempotency_key": "weak-identity"}
    assert _json(app, "POST", f"/api/local/tower-rents/{first_row}/corrections", payload)[0] == 200
    second_batch = _json(app, "POST", "/api/import", {"path": str(sample_workbook)})[1]["batch_id"]
    second_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={second_batch}")[1]["tower_rents"][0]["row_id"]
    detail = _json(app, "GET", f"/api/local/tower-rents/{second_row}?batch_id={second_batch}")[1]
    assert detail["current"]["产品服务费合计（元/年）（不含税）"] == 10000
    assert detail["version"] == 0


def test_duplicate_agreement_in_new_batch_does_not_assign_old_version_arbitrarily(
    app_config, sample_workbook, tmp_path
):
    workbook = load_workbook(sample_workbook)
    sheet = workbook["铁塔租费台账"]
    sheet.cell(1, 9, "需求单号")
    sheet.cell(2, 9, "ORDER-1")
    initial = tmp_path / "rent-initial.xlsx"
    workbook.save(initial)
    initialize_database(app_config)
    app = create_app(app_config)
    first_batch = _json(app, "POST", "/api/import", {"path": str(initial)})[1]["batch_id"]
    first_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={first_batch}")[1]["tower_rents"][0]["row_id"]
    payload = {"batch_id": first_batch,
               "changes": {"产品服务费合计（元/年）（不含税）": 5000},
               "evidence": "核实材料", "operator": "省公司", "error_cause": "计费错误",
               "source": "合同", "idempotency_key": "strong-identity"}
    assert _json(app, "POST", f"/api/local/tower-rents/{first_row}/corrections", payload)[0] == 200
    sheet.append([2, "HZ001", "西湖一站", "杭州", "西湖", "TT001", "另一合同", 6000, "ORDER-1"])
    ambiguous = tmp_path / "rent-ambiguous.xlsx"
    workbook.save(ambiguous)
    second_batch = _json(app, "POST", "/api/import", {"path": str(ambiguous)})[1]["batch_id"]
    rows = _json(app, "GET", f"/api/local/tower-rents?batch_id={second_batch}")[1]["tower_rents"]
    assert [row["current_version"] for row in rows] == [0, 0]
    assert _json(app, "GET", f"/api/local/tower-rents/{first_row}?batch_id={first_batch}")[1]["version"] == 1


def test_archived_batch_keeps_effective_rent_snapshot_after_new_batch_changes(
    app_config, sample_workbook, tmp_path
):
    workbook = load_workbook(sample_workbook)
    sheet = workbook["铁塔租费台账"]
    sheet.cell(1, 9, "需求单号")
    sheet.cell(2, 9, "ORDER-ARCHIVE")
    path = tmp_path / "rent-history.xlsx"
    workbook.save(path)
    initialize_database(app_config)
    app = create_app(app_config)
    first_batch = _json(app, "POST", "/api/import", {"path": str(path)})[1]["batch_id"]
    first_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={first_batch}")[1]["tower_rents"][0]["row_id"]
    payload = {"batch_id": first_batch,
               "changes": {"产品服务费合计（元/年）（不含税）": 5000},
               "evidence": "首批核实", "operator": "省公司", "error_cause": "计费错误",
               "source": "合同", "idempotency_key": "archive-first"}
    assert _json(app, "POST", f"/api/local/tower-rents/{first_row}/corrections", payload)[0] == 200
    with database_for(app_config).unit_of_work() as unit:
        unit.batches.update_status(first_batch, "archived", archive=True)
    second_batch = _json(app, "POST", "/api/import", {"path": str(path)})[1]["batch_id"]
    second_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={second_batch}")[1]["tower_rents"][0]["row_id"]
    assert _json(app, "POST", f"/api/local/tower-rents/{second_row}/corrections",
                 {**payload, "batch_id": second_batch,
                  "changes": {"产品服务费合计（元/年）（不含税）": 4000},
                  "idempotency_key": "archive-second"})[0] == 200
    old_detail = _json(app, "GET", f"/api/local/tower-rents/{first_row}?batch_id={first_batch}")[1]
    new_detail = _json(app, "GET", f"/api/local/tower-rents/{second_row}?batch_id={second_batch}")[1]
    field = "产品服务费合计（元/年）（不含税）"
    assert old_detail["source"][field] == 10000
    assert old_detail["current"][field] == 5000
    assert old_detail["version"] == 1 and len(old_detail["versions"]) == 1
    assert new_detail["current"][field] == 4000
    assert new_detail["version"] == 2
    old_rows = _json(app, "GET", f"/api/ledger-rows?batch_id={first_batch}&ledger_type=tower_rent")[1]["rows"]
    assert old_rows[0]["raw"][field] == 5000


def test_migration_preserves_distinct_old_archived_rent_values(
    app_config, sample_workbook, tmp_path
):
    workbook = load_workbook(sample_workbook)
    sheet = workbook["铁塔租费台账"]
    sheet.cell(1, 9, "需求单号")
    sheet.cell(2, 9, "ORDER-OLD")
    first_path = tmp_path / "old-first.xlsx"
    workbook.save(first_path)
    sheet.cell(2, 8, 6000)
    second_path = tmp_path / "old-second.xlsx"
    workbook.save(second_path)
    initialize_database(app_config)
    app = create_app(app_config)
    first_batch = _json(app, "POST", "/api/import", {"path": str(first_path)})[1]["batch_id"]
    second_batch = _json(app, "POST", "/api/import", {"path": str(second_path)})[1]["batch_id"]
    before_migration_row = _json(app, "GET", f"/api/local/tower-rents?batch_id={second_batch}")[1]["tower_rents"][0]["row_id"]
    before_migration = _json(app, "GET", f"/api/local/tower-rents/{before_migration_row}?batch_id={second_batch}")[1]
    assert before_migration["current"]["产品服务费合计（元/年）（不含税）"] == 6000
    with connect(app_config) as db:
        for table in ("tower_rent_change_requests", "authoritative_tower_rent_versions",
                      "authoritative_tower_rent_sources", "authoritative_tower_rents"):
            db.execute(f"drop table {table}")
        db.execute("update import_batches set is_archived = 1 where id in (?, ?)",
                   (first_batch, second_batch))
        db.execute("delete from schema_migrations where version = 9")
        db.commit()
        apply_migrations(db)
    field = "产品服务费合计（元/年）（不含税）"
    for batch_id, expected in ((first_batch, 10000), (second_batch, 6000)):
        row_id = _json(app, "GET", f"/api/local/tower-rents?batch_id={batch_id}")[1]["tower_rents"][0]["row_id"]
        detail = _json(app, "GET", f"/api/local/tower-rents/{row_id}?batch_id={batch_id}")[1]
        assert detail["source"][field] == expected
        assert detail["current"][field] == expected
        assert detail["version"] == 0


def test_duplicate_rent_business_key_still_keeps_independent_records(
    app_config, sample_workbook, tmp_path
):
    workbook = load_workbook(sample_workbook)
    sheet = workbook["铁塔租费台账"]
    sheet.append([2, "HZ001", "西湖一站", "杭州", "西湖", "TT001", "另一合同", 6000])
    path = tmp_path / "ambiguous-rents.xlsx"
    workbook.save(path)
    initialize_database(app_config)
    app = create_app(app_config)
    batch_id = _json(app, "POST", "/api/import", {"path": str(path)})[1]["batch_id"]
    rows = _json(app, "GET", f"/api/local/tower-rents?batch_id={batch_id}")[1]["tower_rents"]
    first, second = [row["row_id"] for row in rows]
    payload = {"batch_id": batch_id,
               "changes": {"产品服务费合计（元/年）（不含税）": 5000},
               "evidence": "另一合同", "operator": "省公司", "error_cause": "原值错误",
               "source": "核实记录", "idempotency_key": "rent-duplicate"}
    assert _json(app, "POST", f"/api/local/tower-rents/{second}/corrections", payload)[0] == 200
    field = "产品服务费合计（元/年）（不含税）"
    assert _json(app, "GET", f"/api/local/tower-rents/{first}?batch_id={batch_id}")[1]["current"][field] == 10000
    assert _json(app, "GET", f"/api/local/tower-rents/{second}?batch_id={batch_id}")[1]["current"][field] == 5000
    assert _json(app, "POST", f"/api/local/tower-rents/{second}/corrections",
                 {**payload, "changes": {"需求单号": "wrong"},
                  "idempotency_key": "invalid-key"})[0] == 400
    assert _json(app, "GET", f"/api/local/tower-rents/9999?batch_id={batch_id}")[0] == 404


def test_existing_tower_rent_rows_are_backfilled_as_source_and_authority(
    app_config, sample_workbook
):
    initialize_database(app_config)
    app = create_app(app_config)
    batch_id = _json(app, "POST", "/api/import", {"path": str(sample_workbook)})[1]["batch_id"]
    with connect(app_config) as db:
        for table in ("tower_rent_change_requests", "authoritative_tower_rent_versions",
                      "authoritative_tower_rent_sources", "authoritative_tower_rents"):
            db.execute(f"drop table {table}")
        db.execute("delete from schema_migrations where version = 9")
        db.commit()
        apply_migrations(db)
    row = _json(app, "GET", f"/api/local/tower-rents?batch_id={batch_id}")[1]["tower_rents"][0]
    detail = _json(app, "GET", f"/api/local/tower-rents/{row['row_id']}?batch_id={batch_id}")[1]
    assert detail["source"] == detail["current"]
    assert detail["version"] == 0
    assert detail["versions"] == []


def test_local_tower_rent_rejects_invalid_and_archived_edits(app_config, sample_workbook):
    initialize_database(app_config)
    app = create_app(app_config)
    batch_id = _json(app, "POST", "/api/import", {"path": str(sample_workbook)})[1]["batch_id"]
    row_id = _json(app, "GET", f"/api/local/tower-rents?batch_id={batch_id}")[1]["tower_rents"][0]["row_id"]
    path = f"/api/local/tower-rents/{row_id}/corrections"
    payload = {"batch_id": batch_id,
               "changes": {"产品服务费合计（元/年）（不含税）": 5000},
               "evidence": "核实材料", "operator": "省公司", "error_cause": "计费错误",
               "source": "合同", "idempotency_key": "archived-rent"}
    assert _json(app, "POST", path, {**payload, "evidence": ""})[0] == 400
    assert _json(app, "POST", path, {**payload, "batch_id": "invalid"})[0] == 400
    assert app.handle_test_request("POST", path, "{")[0] == 400
    assert _json(app, "POST", path, {**payload, "changes": {"地市": "宁波"}})[0] == 400
    assert _json(app, "POST", path, {**payload, "changes": {"不存在的字段": 1}})[0] == 409
    assert _json(app, "POST", "/api/local/tower-rents/99999/corrections", payload)[0] == 409
    assert _json(app, "GET", f"/api/local/tower-rents/not-a-number?batch_id={batch_id}")[0] == 404
    with connect(app_config) as db:
        db.execute("update import_batches set is_archived = 1 where id = ?", (batch_id,))
    assert _json(app, "POST", path, payload)[0] == 409
